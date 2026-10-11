# -*- coding: utf-8 -*-
import sys
sys.path.insert(0, '/tmp')
from gen_testcases_part1 import ROWS
from gen_testcases_part2 import ROWS2

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ALL = ROWS + ROWS2

wb = Workbook()

# ---------- Sheet 1: 测试用例 ----------
ws = wb.active
ws.title = "测试用例"
headers = ["用例编号", "所属模块", "用例标题", "优先级", "前置条件", "测试步骤", "预期结果", "用例类型", "自动化方式", "执行结果", "备注"]
ws.append(headers)

header_fill = PatternFill("solid", fgColor="1F6FEB")
header_font = Font(color="FFFFFF", bold=True, size=11)
thin = Side(style="thin", color="D0D0D0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
prio_fill = {"P0": "FDE3E3", "P1": "FFF3D6", "P2": "E8F4FD"}
module_fill = PatternFill("solid", fgColor="F3F4F6")

for i, r in enumerate(ALL, start=2):
    ws.append(r + ["", ""])
    ws.cell(row=i, column=4).fill = PatternFill("solid", fgColor=prio_fill.get(r[3], "FFFFFF"))
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=i, column=c)
        cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.cell(row=i, column=1).fill = module_fill

widths = [16, 10, 26, 8, 22, 46, 42, 8, 20, 10, 18]
for idx, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(idx)].width = w
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(ALL)+1}"
ws.row_dimensions[1].height = 24

# ---------- Sheet 2: 统计总览 ----------
st = wb.create_sheet("统计总览")
st.append(["materin-ops 格物 · 测试用例统计"])
st["A1"].font = Font(bold=True, size=14)
st.append([])
st.append(["模块", "用例数", "P0", "P1", "P2"])
for c in range(1, 6):
    cell = st.cell(row=3, column=c)
    cell.fill = header_fill; cell.font = header_font
from collections import Counter
by_mod = {}
for r in ALL:
    m = r[1]
    d = by_mod.setdefault(m, [0, 0, 0, 0])
    d[0] += 1
    d[{"P0": 1, "P1": 2, "P2": 3}[r[3]]] += 1
row = 4
for m, d in by_mod.items():
    st.append([m, d[0], d[1], d[2], d[3]])
    row += 1
st.append(["合计", len(ALL),
           sum(1 for r in ALL if r[3] == "P0"),
           sum(1 for r in ALL if r[3] == "P1"),
           sum(1 for r in ALL if r[3] == "P2")])
st.column_dimensions['A'].width = 14
for c in range(2, 6):
    st.column_dimensions[get_column_letter(c)].width = 10

# ---------- Sheet 3: 自动化映射 ----------
mp = wb.create_sheet("自动化映射")
mp.append(["自动化层", "工具", "位置", "运行方式", "覆盖用例"])
for c in range(1, 6):
    cell = mp.cell(row=1, column=c)
    cell.fill = header_fill; cell.font = header_font
mp.append(["后端单元测试", "JUnit5 + Mockito", "backend/*/src/test/java", "mvn test -pl components/ticket,components/ops -am", "TC-TKT-002~010, TC-REQ-003/004/010"])
mp.append(["前端单元测试", "Vitest", "frontend/apps/web-antd/src/**/__tests__", "根目录 pnpm exec vitest run", "shared 纯函数（labelOf/flattenModules）"])
mp.append(["E2E", "Playwright + Chromium", "e2e/", "cd e2e && npx playwright test", "TC-SYS-001/002/003, TC-TKT-001/007/011"])
mp.append(["API 自动化/压测", "JMeter 5.6.3（Docker）", "jmeter/materin-ops-test.jmx", "见 jmeter/README.md", "TC-TKT 分页/创建, TC-NTC, 登录断言"])
mp.append(["冒烟测试", "bash + curl + python3", "jmeter/smoke.sh", "bash jmeter/smoke.sh", "12 项主链路含 HMAC 签名"])
mp.append(["浏览器验收", "Chrome DevTools MCP", "会话内", "人工触发", "截图留档 .claude/shots/"])
for row in mp.iter_rows(min_row=2):
    for cell in row:
        cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
for idx, w in enumerate([16, 22, 34, 44, 44], start=1):
    mp.column_dimensions[get_column_letter(idx)].width = w

out = "/Users/qiyuyang/developCode/self/materin-ops/docs/materin-ops-测试用例.xlsx"
wb.save(out)
print("Excel saved:", out, "| 用例总数:", len(ALL))
